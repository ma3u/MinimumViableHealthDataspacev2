#!/usr/bin/env python3
"""Builds docs/demos/basque-ehds-15min.pptx, the 15-minute pitch for the Basque
Government's Department of Health.

The Department sent eight questions in advance. This deck answers them in that
order, one slide each, between a title, an overview of the eight answers, and a
closing slide with next steps. Eleven slides, about fifteen minutes; every
slide carries speaker notes with its time budget.

The audience is Spanish-speaking, so the language is simple English: short
sentences, common words, no idioms.

House style is the Spain decks', reused rather than reinvented: navy rule, bold
navy title, grey standfirst, cream cards with a coloured header bar, the same
five accents. The Basque Country photograph comes out of the long Spain deck at
build time, so there is no second copy of it in the repository.

    python3 docs/demos/build-basque-deck.py

Rebuilds from scratch each time, so it is safe to re-run after editing.
"""

import io
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
DIAGRAMS = HERE.parent / "diagrams"
LONG_DECK = HERE / "spain-ehds-ministry-deck.pptx"
OUT = HERE / "basque-ehds-15min.pptx"
LAYERS = DIAGRAMS / "ehds-layers-simple.png"

NAVY = RGBColor(0x14, 0x3D, 0x59)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
AMBER = RGBColor(0xE6, 0xA8, 0x17)
RED = RGBColor(0xC3, 0x1E, 0x2E)
GREEN = RGBColor(0x52, 0x8B, 0x55)
CREAM = RGBColor(0xF4, 0xEF, 0xE6)
BODY = RGBColor(0x4F, 0x55, 0x60)
FAINT = RGBColor(0x8A, 0x8F, 0x98)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FOOTER = "European Health Data Space  ·  Lessons for the Basque Country"

QUESTIONS = [
    "Main challenges",
    "Lessons learned",
    "Decision criteria",
    "Reference architecture",
    "Components and maturity",
    "Governance and operating model",
    "Costs",
    "Effort and team",
]


# ------------------------------------------------------------------- helpers
def shape(slide, kind, x, y, w, h, fill):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.05
    return s


def textbox(slide, x, y, w, h, lines, size, colour=BODY, bold=False, spacing=1.0,
            align=None, anchor=None):
    """Lines are strings, or (lead, rest) pairs where the lead is bold navy."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if anchor is not None:
        tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        if align is not None:
            p.alignment = align
        parts = line if isinstance(line, tuple) else (line,)
        for j, part in enumerate(parts):
            run = p.add_run()
            run.text = part
            run.font.name = "Calibri"
            run.font.size = Pt(size)
            run.font.bold = bold or (len(parts) > 1 and j == 0)
            run.font.color.rgb = NAVY if (len(parts) > 1 and j == 0) else colour
    return tb


def sources(slide, text):
    textbox(slide, 6.4, 7.05, 6.3, 0.3, [text], 9, FAINT, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, header, accent, lines, size=14, head_size=17):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, CREAM)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, 0.45, accent)
    textbox(slide, x + 0.22, y + 0.08, w - 0.44, 0.34, [header], head_size, WHITE, bold=True)
    textbox(slide, x + 0.22, y + 0.65, w - 0.44, h - 0.8, lines, size, BODY, spacing=1.1)


def band(slide, y, h, text, accent=NAVY, size=15):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, y, 12.13, h, CREAM)
    shape(slide, MSO_SHAPE.RECTANGLE, 0.6, y, 0.18, h, accent)
    textbox(slide, 1.05, y, 11.4, h, [text], size, NAVY, bold=True, spacing=1.1,
            anchor=MSO_ANCHOR.MIDDLE)


def pill(slide, x, y, w, h, text, fill, colour=WHITE, size=13):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)
    s.adjustments[0] = 0.5
    textbox(slide, x, y, w, h, [text], size, colour, bold=True, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
    return s


def badge(slide, x, y, d, text, fill, size=16):
    shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill)
    textbox(slide, x, y, d, d, [text], size, WHITE, bold=True, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def question_frame(prs, n, title, standfirst):
    """A slide that answers question n: the question as a label, the answer as title."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.33, 0.18, NAVY)
    label = f"QUESTION {n} OF 8  ·  {QUESTIONS[n - 1].upper()}"
    pill(slide, 0.6, 0.38, 0.6 + 0.085 * len(label), 0.36, label, TEAL, size=11)
    textbox(slide, 0.6, 0.85, 12.1, 0.65, [title], 30, NAVY, bold=True)
    textbox(slide, 0.6, 1.52, 12.1, 0.45, [standfirst], 16, BODY)
    textbox(slide, 0.6, 7.05, 5.6, 0.3, [FOOTER], 10, BODY)
    # Progress: eight small squares top right, the current one filled.
    for i in range(8):
        shape(slide, MSO_SHAPE.RECTANGLE, 10.95 + i * 0.23, 0.47, 0.17, 0.17,
              TEAL if i + 1 == n else CREAM)
    return slide


def deck_image(slide_index: int, pic_index: int = 0) -> io.BytesIO:
    """An image out of the long Spain deck, so there is no second copy in the repo."""
    prs = Presentation(str(LONG_DECK))
    pictures = [sh for sh in prs.slides[slide_index].shapes if sh.shape_type == 13]
    return io.BytesIO(pictures[pic_index].image.blob)


# --------------------------------------------------------------------- slides
def title_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 6.6, 7.5, NAVY)
    # The Basque Country photo from the long deck's "17 regions" slide, cropped
    # to the Guggenheim and without its caption band.
    pic = s.shapes.add_picture(deck_image(6, 3), Inches(6.7), Inches(0), Inches(6.633), Inches(7.5))
    pic.crop_bottom = 0.13
    pic.crop_left = pic.crop_right = 0.28
    shape(s, MSO_SHAPE.RECTANGLE, 6.6, 0, 0.1, 7.5, AMBER)
    textbox(s, 0.6, 1.2, 5.7, 2.2, ["A Connected Health System for the Basque Country"],
            38, WHITE, bold=True, spacing=0.95)
    textbox(s, 0.6, 3.55, 5.7, 0.9,
            ["What we learned from Catena-X, and what it means for the EHDS"], 20, WHITE,
            spacing=1.05)
    textbox(s, 0.6, 4.55, 5.7, 0.4, ["Answers to your eight questions"], 18, AMBER, bold=True)
    shape(s, MSO_SHAPE.RECTANGLE, 0.6, 5.55, 1.2, 0.06, AMBER)
    textbox(s, 0.6, 5.8, 5.7, 0.9, [
        "Basque Government  ·  Department of Health",
        "October 2026",
    ], 14, WHITE, spacing=1.2)
    textbox(s, 0.6, 7.0, 5.7, 0.3, ["Sopra Steria  ·  European Health Data Space practice"],
            10, FAINT)
    notes(s, """[0:30]  Good morning. Egun on. Thank you for the invitation.

You sent us eight questions. I will answer them in your order, one slide each.

I speak from experience in Catena-X, the dataspace of the automotive industry, and from building an EHDS demonstrator that runs live today.""")


def overview_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.33, 0.18, NAVY)
    textbox(s, 0.6, 0.38, 12.0, 0.7, ["Your eight questions, in one sentence each"], 30, NAVY,
            bold=True)
    textbox(s, 0.6, 1.08, 12.0, 0.45, ["The short answers first. Each one has its own slide."],
            16, BODY)
    answers = [
        ("The hardest part was culture, not technology.", RED),
        ("Build on existing standards and communities.", TEAL),
        ("Open standards, decided in public.", NAVY),
        ("The Catena-X operating model.", GREEN),
        ("The protocol is stable. Secure environments are still young.", AMBER),
        ("Keep rule-setting and operation separate.", TEAL),
        ("No licences. You pay for people and hosting.", GREEN),
        ("8 to 12 people for a pilot.", NAVY),
    ]
    for i, (answer, accent) in enumerate(answers):
        col, row = i % 4, i // 4
        x, y = 0.6 + col * 3.07, 1.85 + row * 2.55
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 2.85, 2.3, CREAM)
        shape(s, MSO_SHAPE.RECTANGLE, x, y, 2.85, 0.12, accent)
        badge(s, x + 0.22, y + 0.35, 0.55, str(i + 1), accent)
        textbox(s, x + 0.9, y + 0.38, 1.8, 0.5, [QUESTIONS[i]], 13, accent, bold=True,
                anchor=MSO_ANCHOR.MIDDLE)
        textbox(s, x + 0.22, y + 1.1, 2.45, 1.1, [answer], 15, NAVY, bold=True, spacing=1.05)
    textbox(s, 0.6, 7.05, 5.6, 0.3, [FOOTER], 10, BODY)
    notes(s, """[0:30]  Here are the eight answers in one page. If you remember one slide, remember this one.

I will now go through them one by one.""")


def q1(prs):
    s = question_frame(prs, 1, "The hardest part was culture, not technology",
                       "What we met in Catena-X, and what it means for health.")
    card(s, 0.6, 2.2, 5.95, 3.4, "Culture", RED, [
        "Open source and open collaboration were new for most companies and institutions.",
        "",
        "The committees and working groups delivered only when top management gave its "
        "experts time and a mandate.",
    ], size=16)
    card(s, 6.78, 2.2, 5.95, 3.4, "Meaning of the data", AMBER, [
        "Agreeing what the data means took far more work than moving it.",
        "",
        "In health this means FHIR, terminologies such as SNOMED CT and ICD-10, and OMOP.",
    ], size=16)
    band(s, 5.85, 0.95,
         "Moving data is the easy part. Agreeing on what it means is the hard part.",
         accent=RED, size=18)
    sources(s, "Catena-X Automotive Network  ·  Eclipse Dataspace and Eclipse Tractus-X")
    notes(s, """[1:30]  The biggest challenge was culture, not technology.

Open source was new for most companies. People were used to buying products, not to building together in public.

The working groups delivered only when top management gave its experts time and a clear mandate. Without that, meetings happen, but nothing gets built.

On the technical side: moving data is easy. Agreeing what the data means is hard. In health, that means FHIR, the terminologies, and OMOP.""")


def q2(prs):
    s = question_frame(prs, 2, "Four lessons to avoid common mistakes",
                       "What we would do again, and what we recommend for the Basque Country.")
    lessons = [
        ("1", "Build on what exists", TEAL,
         "Run open working groups. Use FHIR, ICD-10, SNOMED CT, OMOP and HealthDCAT-AP. "
         "Do not invent new standards."),
        ("2", "Start small", GREEN,
         "Choose one or two use cases. Show value that you can measure."),
        ("3", "Give back", AMBER,
         "Contribute improvements to the open-source projects. Do not keep private copies."),
        ("4", "Buy standards, not products", RED,
         "Write the tender around standards and conformity tests. This avoids vendor lock-in."),
    ]
    for i, (n, head, accent, text) in enumerate(lessons):
        x = 0.6 + i * 3.07
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.2, 2.85, 3.95, CREAM)
        shape(s, MSO_SHAPE.RECTANGLE, x, 2.2, 2.85, 0.12, accent)
        badge(s, x + 0.22, 2.55, 0.6, n, accent, size=18)
        textbox(s, x + 0.22, 3.35, 2.45, 0.75, [head], 18, NAVY, bold=True, spacing=1.0)
        textbox(s, x + 0.22, 4.15, 2.45, 1.95, [text], 15.5, BODY, spacing=1.12)
    band(s, 6.35, 0.6, "A private copy of open source becomes your own product to maintain.",
         accent=AMBER, size=15)
    sources(s, "hl7.org/fhir  ·  ohdsi.org  ·  HealthDCAT-AP")
    notes(s, """[1:30]  Four lessons.

One: build on what already exists. FHIR, SNOMED CT, ICD-10, OMOP, HealthDCAT-AP. These communities already did the hard work. Do not invent your own standard.

Two: start small. One or two use cases, with a result you can measure.

Three: give back. If you fix something in an open-source project, send the fix back. A private copy becomes your own product, and you pay for it forever.

Four: in the tender, ask for standards and conformity tests, not for a product name. That is how you avoid vendor lock-in.""")


def q3(prs):
    s = question_frame(prs, 3, "Every decision is written down, and everyone can read it",
                       "How decisions were made, and the criteria behind them.")
    textbox(s, 0.6, 2.15, 12.0, 0.4, ["How we decide"], 18, NAVY, bold=True)
    steps = [
        (TEAL, "Working groups", "propose"),
        (NAVY, "Contributors who build it", "decide the priorities"),
        (GREEN, "An Architecture Decision Record", "is published for everyone"),
    ]
    for i, (accent, who, what) in enumerate(steps):
        x = 0.6 + i * 4.12
        shape(s, MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, x, 2.65, 3.95, 1.25, accent)
        textbox(s, x + 0.45, 2.7, 3.2, 1.15, [who, what], 16, WHITE, bold=True,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.05)
    textbox(s, 0.6, 4.25, 12.0, 0.4, ["Five criteria for every decision"], 18, NAVY, bold=True)
    criteria = [
        ("Open standards", TEAL),
        ("Data sovereignty", NAVY),
        ("Fits the law: EHDS, GDPR, Data Act", GREEN),
        ("Mature technology", AMBER),
        ("A clear exit strategy", RED),
    ]
    x = 0.6
    widths = [2.0, 2.15, 3.55, 2.2, 2.25]
    for (text, accent), w in zip(criteria, widths):
        pill(s, x, 4.8, w, 0.6, text, accent, size=14)
        x += w + (12.13 - sum(widths)) / 4
    band(s, 5.8, 0.95,
         "Our EHDS demonstrator works the same way: 57 decision records, public on GitHub.",
         accent=GREEN, size=16)
    sources(s, "github.com/ma3u/MinimumViableHealthDataspacev2  ·  docs/ADRs")
    notes(s, """[1:30]  Every decision was written down as an Architecture Decision Record, an ADR. Everyone in the community could read it, and see why.

The working groups made the proposals. But in the end, the people who build the software decided the priorities. That is how open source works.

Five criteria: open standards, data sovereignty, fit with the law (EHDS, GDPR and the Data Act), mature technology, and a clear exit strategy.

The exit strategy is important for a public body: you must always be able to change the supplier.

Our demonstrator follows the same rule. It has 57 public decision records.""")


def q4(prs):
    s = question_frame(prs, 4, "The reference architecture: the Catena-X operating model",
                       "A service map with clear roles. The EHDS defines the roles, but not who runs them.")
    s.shapes.add_picture(str(LAYERS), Inches(0.6), Inches(2.05), Inches(12.13), Inches(4.04))
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 6.15, 5.95, 0.8, CREAM)
    shape(s, MSO_SHAPE.RECTANGLE, 0.6, 6.15, 0.15, 0.8, GREEN)
    textbox(s, 0.85, 6.2, 5.5, 0.7, [("Implemented:  ", "the Catena-X operating model, with a "
                                                       "service map and clear responsibilities.")],
            13, BODY, spacing=1.05, anchor=MSO_ANCHOR.MIDDLE)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.78, 6.15, 5.95, 0.8, CREAM)
    shape(s, MSO_SHAPE.RECTANGLE, 6.78, 6.15, 0.15, 0.8, AMBER)
    textbox(s, 7.03, 6.2, 5.5, 0.7, [("Considered:  ", "the EHDS role model alone. It has no "
                                                       "answer for operations and onboarding.")],
            13, BODY, spacing=1.05, anchor=MSO_ANCHOR.MIDDLE)
    sources(s, "catenax-ev.github.io/docs/operating-model  ·  Regulation (EU) 2025/327, Chapter IV")
    notes(s, """[1:45]  We implemented the Catena-X operating model. It has a service map: a list of every service, with a clear owner and clear responsibilities.

The EHDS regulation already gives us a role model: the health data access body, the data holders, the data users, and the secure processing environments.

But the regulation does not say who runs the system day to day. Who connects a new hospital? Who answers the phone when something breaks? That is the orange layer, the operating company.

So the EHDS role model alone was not enough. We added the Catena-X operating model below it.""")


def q5(prs):
    s = question_frame(prs, 5, "The components, and how mature they are",
                       "The components follow the service map. Health standards are mature; "
                       "secure environments are not yet.")
    card(s, 0.6, 2.2, 3.9, 4.6, "The service map", NAVY, [], size=1)
    for i, comp in enumerate(["Connectors", "Identity and credentials", "Catalogue",
                              "Onboarding", "Conformity testing"]):
        pill(s, 0.85, 2.9 + i * 0.75, 3.4, 0.55, comp, WHITE, NAVY, size=14)
    rows = [
        (GREEN, "Stable", "Dataspace Protocol",
         "Stable. On its way to an ISO standard, ISO/IEC 26450."),
        (GREEN, "Mature", "FHIR and OMOP",
         "Mature health standards, used around the world."),
        (AMBER, "Changing", "Tractus-X connector",
         "In production since 2023, but still version 0.x. Breaking changes every quarter."),
        (RED, "Young", "Secure processing and data permits",
         "The least mature parts. Plan time and budget here."),
    ]
    for i, (accent, level, name, text) in enumerate(rows):
        y = 2.2 + i * 1.18
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 4.75, y, 7.98, 1.05, CREAM)
        pill(s, 4.95, y + 0.27, 1.45, 0.5, level, accent, size=13)
        textbox(s, 6.6, y + 0.12, 6.0, 0.4, [name], 16, NAVY, bold=True)
        textbox(s, 6.6, y + 0.5, 6.0, 0.5, [text], 13, BODY)
    sources(s, "eclipse-dataspace-protocol-base.github.io  ·  eclipse-tractusx.github.io  ·  hl7.org/fhir  ·  ohdsi.org")
    notes(s, """[1:45]  On the left, the service map: connectors, identity and credentials, the catalogue, onboarding, and conformity testing.

On the right, how mature they are.

The Dataspace Protocol is stable. It is on its way to becoming an ISO standard.

FHIR and OMOP are mature.

The Tractus-X connector has run in production since 2023. But it is still version zero-point-something, and it has breaking changes every quarter. You need a team that can keep up.

The least mature parts are the secure processing environments and the data permit workflows. This is where to plan extra time and budget.""")


def q6(prs):
    s = question_frame(prs, 6, "Governance: three roles, kept apart",
                       "Catena-X separates the rules, the code and the operation. "
                       "We suggest the same for the Basque Country.")
    roles = [
        (NAVY, "Rules", "Catena-X Association", "Sets the standards and runs certification.",
         "Health Data Office", "acts as the health data access body."),
        (TEAL, "Code", "Eclipse Foundation", "Hosts the open-source code.",
         "Open-source projects", "use them, and contribute back."),
        (AMBER, "Operation", "Cofinity-X", "Runs the operation and the core services.",
         "EJIE", "operates the core services."),
    ]
    textbox(s, 0.6, 2.15, 2.0, 0.4, ["Catena-X"], 14, FAINT, bold=True)
    textbox(s, 0.6, 4.45, 3.0, 0.4, ["Basque Country"], 14, FAINT, bold=True)
    for i, (accent, role, cx, cx_text, bq, bq_text) in enumerate(roles):
        x = 0.6 + i * 4.12
        card(s, x, 2.5, 3.9, 1.7, f"{role}:  {cx}", accent, [cx_text], size=14, head_size=15)
        arrow = shape(s, MSO_SHAPE.DOWN_ARROW, x + 1.7, 4.25, 0.5, 0.45, accent)
        arrow.adjustments[0] = 0.5
        sh = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 4.8, 3.9, 1.05, CREAM)
        sh.line.color.rgb = accent
        sh.line.width = Pt(2.25)
        textbox(s, x + 0.22, 4.9, 3.5, 0.85, [bq, bq_text], 15, BODY, spacing=1.05,
                anchor=MSO_ANCHOR.MIDDLE)
        s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
        s.shapes[-1].text_frame.paragraphs[0].runs[0].font.color.rgb = NAVY
    band(s, 6.05, 0.85,
         "Keep rule-setting and operation separate. Design the governance for several "
         "domains from the start, not only for health.", accent=TEAL, size=15)
    sources(s, "catena-x.net  ·  eclipse.org  ·  cofinity-x.com")
    notes(s, """[1:45]  Catena-X keeps three functions separate.

The association sets the standards and runs certification. The Eclipse Foundation hosts the open-source code. And Cofinity-X, a company, runs the operation.

For the Basque Country, this could look like this: the Health Data Office acts as the data access body. It sets the rules and decides on data permits. EJIE operates the core services. And both use, and contribute to, the same open-source projects.

Our two recommendations: keep rule-setting and operation in different hands. And design the governance for more than one domain from the start. Health is the first dataspace, but it will not be the last: energy, mobility, public administration.""")


def q7(prs):
    s = question_frame(prs, 7, "Costs: no licences, but people and hosting",
                       "The core stack is open source. The money goes into engineering, hosting "
                       "and operations.")
    textbox(s, 0.6, 2.1, 6.0, 0.4, ["Catena-X, for reference"], 14, FAINT, bold=True)
    ref = [
        (GREEN, "≈ €0", "Licence costs", "The core stack is open source."),
        (NAVY, "> €100M", "Public funding", "For a whole industry, not one region."),
        (AMBER, "€1,500 – 71,000", "Per participant, per year", "Paid to the operator."),
    ]
    for i, (accent, big, head, text) in enumerate(ref):
        x = 0.6 + i * 4.12
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.5, 3.9, 1.85, CREAM)
        shape(s, MSO_SHAPE.RECTANGLE, x, 2.5, 0.12, 1.85, accent)
        textbox(s, x + 0.35, 2.62, 3.4, 0.7, [big], 30, accent, bold=True)
        textbox(s, x + 0.35, 3.35, 3.4, 0.4, [head], 15, NAVY, bold=True)
        textbox(s, x + 0.35, 3.75, 3.4, 0.5, [text], 13, BODY)
    textbox(s, 0.6, 4.6, 6.0, 0.4, ["Basque regional pilot, our estimate"], 14, FAINT, bold=True)
    est = [
        (TEAL, "€0.8 – 1.5M", "To build the pilot", "Engineering and onboarding."),
        (RED, "15 – 25% a year", "To operate it", "Of the build cost: about €120,000 – 375,000 a year."),
        (AMBER, "€3,000 a month", "Hosting: prod, test and dev",
         "€1,000 per environment, €36,000 a year. Plus SDLC and developer tools."),
    ]
    for i, (accent, big, head, text) in enumerate(est):
        x = 0.6 + i * 4.12
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 5.0, 3.9, 1.85, NAVY)
        shape(s, MSO_SHAPE.RECTANGLE, x, 5.0, 0.12, 1.85, accent)
        textbox(s, x + 0.35, 5.1, 3.4, 0.6, [big], 26, WHITE, bold=True)
        textbox(s, x + 0.35, 5.72, 3.4, 0.35, [head], 15, AMBER, bold=True)
        textbox(s, x + 0.35, 6.08, 3.4, 0.7, [text], 12, WHITE, spacing=1.05)
    sources(s, "catena-x.net  ·  cofinity-x.com  ·  operations: share of the build cost per year")
    notes(s, """[1:45]  Licence costs are close to zero, because the core stack is open source. The real costs are engineering, hosting and operations.

For reference: Catena-X received more than 100 million euros of public funding. But that was for a whole industry. Each participant pays the operator between 1,500 and 71,000 euros a year.

For a regional pilot in the Basque Country we estimate 0.8 to 1.5 million euros to build. Operations then cost about 15 to 25 percent of that every year. That is roughly 120,000 to 375,000 euros a year.

For hosting we plan three environments: production, test and development, at about 1,000 euros a month each. That is 3,000 euros a month, or 36,000 euros a year. On top of that come the SDLC and developer tools: CI/CD, repositories, testing and security scanning.""")


def q8(prs):
    s = question_frame(prs, 8, "The team: 8 to 12 people for a pilot",
                       "Full-time people, with the right mix of skills.")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 2.2, 3.9, 4.6, NAVY)
    textbox(s, 0.6, 2.75, 3.9, 1.3, ["8 – 12"], 72, WHITE, bold=True, align=PP_ALIGN.CENTER)
    textbox(s, 0.6, 4.15, 3.9, 0.5, ["full-time people"], 22, AMBER, bold=True,
            align=PP_ALIGN.CENTER)
    textbox(s, 0.9, 5.0, 3.3, 1.5, ["Experts need time and a mandate from top management. "
                                    "See question 1."], 14, WHITE, align=PP_ALIGN.CENTER,
            spacing=1.1)
    roles = [
        (NAVY, "Architect", "One person who holds the whole picture."),
        (TEAL, "Connector engineers", "Run and update the dataspace connectors."),
        (GREEN, "Platform and security (DevSecOps)", "Hosting, identity, monitoring, security."),
        (AMBER, "FHIR and OMOP specialist", "Makes the data mean the same thing everywhere."),
        (RED, "Data protection and legal", "GDPR, EHDS, data permits and contracts."),
    ]
    for i, (accent, head, text) in enumerate(roles):
        y = 2.2 + i * 0.94
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 4.75, y, 7.98, 0.82, CREAM)
        badge(s, 4.92, y + 0.14, 0.54, str(i + 1), accent, size=15)
        textbox(s, 5.7, y + 0.06, 6.9, 0.38, [head], 16, NAVY, bold=True)
        textbox(s, 5.7, y + 0.43, 6.9, 0.35, [text], 13, BODY)
    notes(s, """[1:15]  A pilot needs about 8 to 12 full-time people.

An architect, who holds the whole picture. Connector engineers. Platform and security engineers, DevSecOps. A FHIR and OMOP specialist, for the meaning of the data. And data protection and legal expertise, because in health, every data flow needs a legal basis.

And remember question one: these people need time and a mandate from top management.""")


def closing_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.33, 0.18, NAVY)
    textbox(s, 0.6, 0.38, 7.4, 0.7, ["Suggested next steps"], 30, NAVY, bold=True)
    textbox(s, 0.6, 1.08, 7.4, 0.45, ["Small steps, each one useful on its own."], 16, BODY)
    steps = [
        (TEAL, "Choose one or two use cases", "with a result you can measure."),
        (NAVY, "Agree the roles", "the Health Data Office as access body, EJIE as operator."),
        (GREEN, "Start open working groups", "with the hospitals, on existing standards."),
        (AMBER, "Build a pilot", "tendered on standards and conformity tests."),
    ]
    for i, (accent, head, text) in enumerate(steps):
        y = 1.8 + i * 1.05
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, y, 7.4, 0.9, CREAM)
        badge(s, 0.8, y + 0.17, 0.56, str(i + 1), accent, size=16)
        textbox(s, 1.6, y + 0.1, 6.2, 0.4, [head], 17, NAVY, bold=True)
        textbox(s, 1.6, y + 0.5, 6.2, 0.35, [text], 13, BODY)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 6.1, 7.4, 0.8, CREAM)
    shape(s, MSO_SHAPE.RECTANGLE, 0.6, 6.1, 0.15, 0.8, RED)
    textbox(s, 0.95, 6.1, 6.9, 0.8, [
        ("March 2027  ", "access bodies designated."),
        ("March 2029  ", "secondary use applies in full."),
    ], 14, BODY, anchor=MSO_ANCHOR.MIDDLE)
    shape(s, MSO_SHAPE.RECTANGLE, 8.4, 0.18, 4.93, 7.32, NAVY)
    shape(s, MSO_SHAPE.RECTANGLE, 8.4, 0.18, 0.1, 7.32, AMBER)
    textbox(s, 8.9, 1.5, 4.1, 2.4, ["Eskerrik asko", "Gracias", "Thank you"], 34, WHITE,
            bold=True, spacing=1.15)
    shape(s, MSO_SHAPE.RECTANGLE, 8.9, 4.2, 1.2, 0.06, AMBER)
    textbox(s, 8.9, 4.45, 4.1, 1.4, [
        "Live demonstration",
        "ehds.mabu.red",
    ], 16, WHITE, spacing=1.2)
    s.shapes[-1].text_frame.paragraphs[1].runs[0].font.color.rgb = AMBER
    s.shapes[-1].text_frame.paragraphs[1].runs[0].font.bold = True
    textbox(s, 8.9, 6.6, 4.2, 0.5, ["Sopra Steria  ·  European Health Data Space practice"],
            10, FAINT)
    textbox(s, 0.6, 7.05, 5.6, 0.3, [FOOTER], 10, BODY)
    textbox(s, 4.4, 7.05, 3.6, 0.3, ["Regulation (EU) 2025/327, Art. 55(6), 105"], 9, FAINT,
            align=PP_ALIGN.RIGHT)
    notes(s, """[1:00]  What could come next? Four small steps.

One: choose one or two use cases with a result you can measure.
Two: agree the roles. The Health Data Office as access body, EJIE as operator.
Three: start open working groups with the hospitals, on existing standards.
Four: build a pilot, and write the tender around standards and conformity tests.

The dates are fixed by the regulation: access bodies by March 2027, and secondary use in full by March 2029. That is closer than it looks.

Eskerrik asko. Thank you. I am happy to take your questions, or to show the live demonstration.""")


def main() -> int:
    for path in (LONG_DECK, LAYERS):
        if not path.exists():
            print(f"missing {path}")
            return 1
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = "A Connected Health System for the Basque Country"
    prs.core_properties.subject = "European Health Data Space, answers to eight questions"
    prs.core_properties.author = "Matthias Buchhorn"
    for fn in (title_slide, overview_slide, q1, q2, q3, q4, q5, q6, q7, q8, closing_slide):
        fn(prs)
    prs.save(str(OUT))
    print(f"{OUT.name}: {len(Presentation(str(OUT)).slides)} slides, "
          f"{OUT.stat().st_size // 1024} KB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
